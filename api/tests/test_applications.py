from datetime import date, datetime, timedelta
from io import BytesIO

from db.models import Application, Offer
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

HEADER = ["N°", "Organisation", "Type", "Ville", "Fiabilité de la piste", "CV à envoyer",
          "Contact / lien", "Poste visé", "Pourquoi cette piste", "Statut", "Date d'envoi",
          "Relance prévue", "Notes"]


def _row(n: object, org: object, status: str = "À contacter", sent: object = None,
         follow_up: object = None) -> list[object]:
    formula = f'=IF(AND(K{n}<>"",J{n}="Envoyé"),K{n}+10,"")'
    return [n, org, "ONG", "Ville", "Piste non vérifiée", "NP5 FR", "https://ex.org/contact",
            "Chargé de suivi", "raison", status, sent,
            formula if follow_up is None else follow_up, None]


def _workbook() -> bytes:
    """Fichier factice : mêmes en-têtes que le vrai, données inventées."""
    wb = Workbook()
    mdg = wb.active
    assert mdg is not None
    mdg.title = "Madagascar"
    mdg.append(HEADER)
    mdg.append(_row(1, "Org A", "Envoyé", datetime(2026, 10, 20)))  # formule non calculée
    mdg.append(_row(2, "Org B", "Envoyé", datetime(2026, 10, 21), datetime(2026, 11, 2)))
    mdg.append(_row(3, "Org C"))
    mdg.append(_row(4, None))                               # organisation vide
    mdg.append(_row(5, "Org E", "Peut-être"))               # statut inconnu
    mdg.append(_row("x", "Org F"))                          # N° invalide
    mdg.append([None] * len(HEADER))                        # ligne vide : ignorée sans bruit
    lille = wb.create_sheet("Lille")
    lille.append(HEADER)
    lille.append(_row(1, "Org L", "Refusé", "21/10/2026"))
    wb.create_sheet("Synthèse").append(["à ne pas lire"])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _upload(client: TestClient, content: bytes) -> dict:
    files = {"file": ("prospection.xlsx", content,
                      "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    r = client.post("/applications/import-xlsx", files=files)
    assert r.status_code == 200, r.text
    return r.json()


def test_import_maps_columns_and_reports_skipped(client: TestClient, session: Session) -> None:
    report = _upload(client, _workbook())
    assert report["imported"] == 4 and report["already_present"] == 0
    assert [(s["sheet"], s["row"]) for s in report["skipped"]] == [
        ("Madagascar", 5), ("Madagascar", 6), ("Madagascar", 7)]
    assert "statut inconnu" in report["skipped"][1]["reason"]
    apps = {(a.sheet, a.sheet_row): a for a in session.scalars(select(Application))}
    a1 = apps[("Madagascar", 1)]
    assert (a1.organization, a1.position, a1.cv_sent, a1.contact, a1.status) == (
        "Org A", "Chargé de suivi", "NP5 FR", "https://ex.org/contact", "applied")
    assert a1.sent_at == date(2026, 10, 20) and a1.follow_up_at == date(2026, 10, 30)
    assert apps[("Madagascar", 2)].follow_up_at == date(2026, 11, 2)  # valeur calculée lue
    assert apps[("Madagascar", 3)].status == "interesting" and apps[("Madagascar", 3)].follow_up_at is None
    assert apps[("Lille", 1)].status == "rejected" and apps[("Lille", 1)].sent_at == date(2026, 10, 21)


def test_reimport_is_idempotent_and_keeps_interface_changes(client: TestClient,
                                                            session: Session) -> None:
    _upload(client, _workbook())
    app_id = session.scalars(select(Application.id).where(Application.organization == "Org C")).one()
    client.patch(f"/applications/{app_id}", json={"status": "applied", "notes": "envoyée"})
    report = _upload(client, _workbook())
    assert report["imported"] == 0 and report["already_present"] == 4
    assert len(client.get("/applications").json()) == 4
    assert client.get("/applications", params={"status": "applied"}).json()[-1]["notes"] == "envoyée"


def test_import_rejects_non_excel(client: TestClient, session: Session) -> None:
    r = client.post("/applications/import-xlsx", files={"file": ("x.xlsx", b"pas un xlsx")})
    assert r.status_code == 422


def test_create_update_and_due_follow_ups(client: TestClient, session: Session) -> None:
    offer = Offer(title="Poste", canonical_url="https://ex.org/1")
    session.add(offer)
    session.commit()
    past = (date.today() - timedelta(days=3)).isoformat()
    r = client.post("/applications", json={"organization": "Org", "offer_id": offer.id,
                                           "status": "applied", "follow_up_at": past})
    assert r.status_code == 201 and r.json()["follow_up_due"] is True
    app_id = r.json()["id"]
    client.post("/applications", json={"organization": "Plus tard", "status": "applied",
                                       "follow_up_at": (date.today() + timedelta(days=5)).isoformat()})
    assert [a["id"] for a in client.get("/applications", params={"due": True}).json()] == [app_id]
    r = client.patch(f"/applications/{app_id}", json={"status": "interview"})
    assert r.json()["follow_up_due"] is False and r.json()["organization"] == "Org"
    assert client.get("/applications", params={"due": True}).json() == []


def test_application_errors(client: TestClient, session: Session) -> None:
    assert client.post("/applications", json={"status": "envoyé"}).status_code == 422
    assert client.post("/applications", json={"offer_id": 999}).status_code == 422
    assert client.patch("/applications/999", json={"notes": "x"}).status_code == 404
    app_id = client.post("/applications", json={"organization": "Org"}).json()["id"]
    assert client.patch(f"/applications/{app_id}", json={"status": None}).status_code == 422

